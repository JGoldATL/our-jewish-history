#!/usr/bin/env python3
"""Tests for tools/sheet-to-history.py. Builds a tiny workbook (same headers as the master Sheet) with
planted internal values, runs the loader on a scratch copy of history.json, and checks the rules.

    python3 tests/load.py          (needs: pip install openpyxl)
"""
import json, subprocess, sys, tempfile
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SCRIPT = ROOT / 'tools' / 'sheet-to-history.py'
SECRET = 'SECRET-INTERNAL'

HEAD = {
    'Events': ['Event ID', 'Start year', 'End year', 'Date (display)', 'Place ID', 'Movement record IDs', 'Evidence object IDs',
               'Card type', 'Card date', 'Card title', 'Card description', 'Meanwhile (card line)', 'Meanwhile status',
               'Map symbol', 'Card status', 'Internal notes & qualifications (never on cards)', 'Source IDs',
               'Evidence category', 'Confidence', 'Note for Jeffrey', 'Story / what happened', 'Population killed'],
    'Communities': ['Record ID', 'Family', 'Date (display)', 'Start year', 'End year', 'Place ID', 'Map symbol', 'Linked records',
                    'Card type', 'Card date', 'Card title', 'Card description', 'Card status', 'Map treatment',
                    'Internal notes & qualifications (never on cards)', 'Source'],
    'Movements': ['Movement ID', 'Event ID', 'Date (display)', 'Start year', 'End year', 'Origin (historical)',
                  'Destination (historical)', 'Arrow treatment', 'Linked records', 'Internal notes & qualifications (never on cards)',
                  'Population low', 'Source'],
    'Context & Voyages': ['Record ID', 'Family', 'Date (display)', 'Start year', 'Title', 'What it is', 'Linked records', 'Source'],
    'Evidence': ['Evidence ID', 'Creation start yr', 'Creation end yr', 'Linked event IDs', 'Name', 'Internal confidence notes (never on cards)'],
    'Places': ['Place ID', 'Historical name', 'Modern equivalent', 'Role', 'Latitude', 'Longitude', 'Land of Israel?', 'Notes'],
}
S = SECRET
BASE = {
    'Places': [['place-a', 'A', 'a', 'region', 31.0, 35.0, 'Yes', S], ['place-b-c', 'B', 'b', 'city', 32.0, 36.0, 'No', S]],
    'Events': [
        ['EVT-0001', -700, None, '700 BCE', 'place-a', 'MOV-001', None, 'Event', '700 BCE', 'Approved one', 'Plain text one.', 'Meanwhile approved.', 'Approved', None, 'approved', S, S, S, S, S, S, S],
        ['EVT-0002', -600, None, '600 BCE', 'place-b-c', None, None, 'Event', '600 BCE', 'Draft one', 'Plain text two.', 'Meanwhile draft.', 'Draft', None, 'draft', S, S, S, S, S, S, S],
        ['EVT-0003', -500, None, '500 BCE', 'place-a', None, None, 'Archaeology', '500 BCE', 'Arch one', 'Plain text three.', 'Meanwhile on arch.', 'Draft', None, 'draft', S, S, S, S, S, S, S],
        ['EVT-0004', -400, None, '400 BCE', 'place-a', None, None, 'Event', '400 BCE', 'Half card', None, None, None, None, None, S, S, S, S, S, S, S],
        ['EVT-0005', -300, None, '300 BCE', 'place-a', None, None, None, None, None, None, None, None, None, None, S, S, S, S, S, S, S],
    ],
    'Communities': [
        ['POP-001', 'POP', 'c. 1000 BCE', -1000, None, 'place-a', 'Blue presence', 'EVT-0001', None, None, None, None, None, S, S, S],
        ['DES-001', 'DES', '600 BCE', -600, None, 'place-b-c', 'Red octagon', None, 'Event', '600 BCE', 'Draft one', 'Plain text two.', 'approved', S, S, S],
        ['DES-002', 'DES', '100 CE', 100, None, None, 'Red octagon', None, None, None, None, None, None, S, S, S],
        ['CHG-001', 'CHG', '200 CE', 200, 150, 'place-a', None, None, None, None, None, None, None, S, S, S],
    ],
    'Movements': [
        ['MOV-001', 'EVT-0001', '700 BCE', -700, -690, 'A', 'B', 'Red', 'POP-001', S, S, S],
        ['MOV-002', 'EVT-0002', '600 BCE', -600, None, 'A', 'B', 'Not stated in CEO chat', None, S, S, S],
    ],
    'Context & Voyages': [['CTX-001', 'CTX', '500 BCE', -500, S, S, 'POP-001', S]],
    'Evidence': [['EVD-001', -700, -700, 'EVT-0001', S, S]],
}


def make(path, mutate=None):
    data = {k: [list(r) for r in v] for k, v in BASE.items()}
    if mutate:
        mutate(data)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for tab, head in HEAD.items():
        ws = wb.create_sheet(tab); ws.append(head)
        for r in data[tab]: ws.append(r)
    wb.save(path)


def run(xlsx, hist, *extra):
    p = subprocess.run([sys.executable, str(SCRIPT), str(xlsx), '--history', str(hist), '--report', str(hist) + '.report.json', *extra],
                       capture_output=True, text=True)
    rep = json.loads(Path(str(hist) + '.report.json').read_text()) if Path(str(hist) + '.report.json').exists() else None
    return p, rep


fails = []
def check(name, cond, detail=''):
    print(('PASS ' if cond else 'FAIL ') + name + (f'  [{detail}]' if detail and not cond else ''))
    if not cond: fails.append(name)


with tempfile.TemporaryDirectory() as td:
    td = Path(td)
    engine = {'schemaVersion': '1.0', 'places': [{'id': 'place-a', 'x': 1}], 'presentation': {'version': 'keep me'}}
    hist = td / 'history.json'
    hist.write_text(json.dumps(engine, indent=2, ensure_ascii=False), encoding='utf-8')
    x1 = td / 'one.xlsx'; make(x1)

    p, rep = run(x1, hist)
    out = json.loads(hist.read_text())
    sheet = out['sheet']
    check('loader runs', p.returncode == 0, p.stderr[-300:])
    check('engine keys untouched', all(out[k] == engine[k] for k in engine))
    check('no internal text leaks into history.json', SECRET not in hist.read_text())
    check('no forbidden field names in output', not any(w in hist.read_text() for w in ('Population', 'sourceIds', 'confidence', 'evidenceCategory', 'Map treatment')))
    cards = {c['id']: c for c in sheet['cards']}
    check('all complete cards loaded, drafts included in preview', set(cards) == {'EVT-0001', 'EVT-0002', 'EVT-0003'}, list(cards))
    check('record with no card is not in the card list', 'EVT-0005' not in cards and any(e['id'] == 'EVT-0005' and not e['hasCard'] for e in sheet['events']))
    check('incomplete card rejected with reason, record kept', any(r['id'] == 'EVT-0004' for r in rep['rejects']) and any(e['id'] == 'EVT-0004' for e in sheet['events']))
    check('duplicate card text makes one card only', 'DES-001' not in cards and any(w['id'] == 'DES-001' for w in rep['warnings']))
    check('record with no Place ID rejected, not fixed', any(r['id'] == 'DES-002' for r in rep['rejects']) and not any(c['id'] == 'DES-002' for c in sheet['communities']))
    check('end before start rejected', any(r['id'] == 'CHG-001' for r in rep['rejects']))
    check('Meanwhile kept where filled, with its status', cards['EVT-0001']['meanwhile'] == 'Meanwhile approved.' and cards['EVT-0002']['meanwhileStatus'] == 'draft')
    check('Meanwhile on a non-Event card is withheld from the file and reported',
          'meanwhile' not in cards['EVT-0003'] and 'Meanwhile on arch' not in hist.read_text()
          and [w['id'] for w in rep['meanwhileWithheldNonEvent']] == ['EVT-0003'])
    check('arrow "not stated" is flagged', any(w['id'] == 'MOV-002' for w in rep['warnings']))
    check('hyphenated place IDs load', any(p_['id'] == 'place-b-c' for p_ in sheet['places']))
    check('POP/CHG with no engine fields are counted, not guessed', rep['recordsWithNoEngineFields'] >= 1 and not any('engine' in c for c in sheet['communities']))
    check('first load reports everything as added', rep['diff']['cards']['added'] == ['EVT-0001', 'EVT-0002', 'EVT-0003'])

    before = hist.read_text()
    p, rep = run(x1, hist)
    check('second run on the same export changes nothing', hist.read_text() == before and 'No changes' in p.stdout)

    x2 = td / 'two.xlsx'
    def edit(d):
        d['Events'][0][9] = 'Approved one, reworded'       # one card title
        d['Events'].pop(2)                                   # remove EVT-0003
        d['Movements'].append(['MOV-003', None, '1 CE', 1, None, 'A', 'B', 'Blue', None, S, S, S])
    make(x2, edit)
    p, rep = run(x2, hist)
    d = rep['diff']
    check('edit touches only the changed record', d['cards']['changed'] == ['EVT-0001'] and d['events']['changed'] == [], d['cards'])
    check('a reworded title is reported as exactly the title field changing', d['cards']['changedFields'] == {'EVT-0001': ['title']}, d['cards']['changedFields'])
    check('removed record is reported', d['cards']['removed'] == ['EVT-0003'] and d['events']['removed'] == ['EVT-0003'])
    check('added record is reported', d['movements']['added'] == ['MOV-003'])
    check('unchanged records not listed as changed', d['places']['changed'] == [] and d['communities']['changed'] == [])

    x3 = td / 'three.xlsx'
    make(x3)
    pub = td / 'public.json'
    p, rep = run(x3, hist, '--public', '--out', str(pub))
    pj = json.loads(pub.read_text())['sheet']
    pcards = {c['id']: c for c in pj['cards']}
    check('public build: approved cards only', set(pcards) == {'EVT-0001'}, list(pcards))
    check('public build: no draft Meanwhile lines', all('meanwhile' not in c or c.get('meanwhileStatus') == 'approved' for c in pj['cards']))
    check('public build: approved Meanwhile kept', pcards['EVT-0001'].get('meanwhile') == 'Meanwhile approved.')
    check('public file leaks nothing', SECRET not in pub.read_text())
    check('--public --out leaves the preview file alone', json.loads(hist.read_text())['sheet']['meta']['mode'] == 'preview')

    x6 = td / 'six.xlsx'
    make(x6, lambda dd: dd['Events'].append(['Pasted paragraph of chat text that is far longer than any ID should ever be, in column A'] + [None] * 21))
    p, rep = run(x6, hist)
    check('stray text in an ID column is rejected with a reason, not a crash',
          p.returncode == 0 and any(r['tab'] == 'Events' and 'permanent-ID' in r['reason'] for r in rep['rejects']), p.stderr[-200:])

    x7 = td / 'seven.xlsx'
    wb = openpyxl.load_workbook(x1); ws = wb['Movements']
    c0 = ws.max_column
    ws.cell(1, c0 + 1, 'Origin place ID'); ws.cell(1, c0 + 2, 'Destination place ID')
    ws.cell(2, c0 + 1, 'place-a'); ws.cell(2, c0 + 2, 'place-b-c')      # MOV-001: both ends valid
    ws.cell(3, c0 + 1, 'place-a')                                        # MOV-002: origin only
    wb.save(x7)
    p, rep = run(x7, hist)
    mv = {m['id']: m for m in json.loads(hist.read_text())['sheet']['movements']}
    check('Place ID columns load and mark an arrow as ready only when both ends are present',
          mv['MOV-001']['arrowEndsReady'] is True and mv['MOV-002']['arrowEndsReady'] is False and rep['movementsWithArrowEnds'] == 1, mv)
    check('a movement with only one end gets a warning, not a guess',
          any(w['id'] == 'MOV-002' and 'only one end' in w['note'] for w in rep['warnings']) and mv['MOV-002']['destinationPlaceId'] is None)

    x8 = td / 'eight.xlsx'
    wb = openpyxl.load_workbook(x7); ws = wb['Movements']; ws.cell(2, c0 + 2, 'place-nowhere'); wb.save(x8)
    p, rep = run(x8, hist)
    mv = {m['id'] for m in json.loads(hist.read_text())['sheet']['movements']}
    check('an unknown Place ID rejects the movement with the reason',
          any(r['id'] == 'MOV-001' and 'not in the Places tab' in r['reason'] for r in rep['rejects']) and 'MOV-001' not in mv)

    x9 = td / 'nine.xlsx'
    wb = openpyxl.load_workbook(x7); wsP = wb['Places']
    wsP.cell(3, 5).value = None; wsP.cell(3, 6).value = None            # place-b-c: no coordinates yet
    wb.save(x9)
    p, rep = run(x9, hist)
    d9 = json.loads(hist.read_text())['sheet']; pl = {q['id']: q for q in d9['places']}; mv = {m['id']: m for m in d9['movements']}
    check('a Place with no coordinates loads with a warning and is not guessed',
          p.returncode == 0 and 'place-b-c' in pl and pl['place-b-c']['lat'] is None
          and any(w['id'] == 'place-b-c' and 'no latitude/longitude' in w['note'] for w in rep['warnings']))
    check('a movement to a Place with no coordinates is not marked ready to draw',
          'MOV-001' in mv and mv['MOV-001']['arrowEndsReady'] is False and rep['movementsWithArrowEnds'] == 0)
    x10 = td / 'ten.xlsx'
    wb = openpyxl.load_workbook(x7); wb['Places'].cell(3, 5).value = 95.0; wb['Places'].cell(3, 6).value = None; wb.save(x10)
    p, rep = run(x10, hist)
    check('a Place with a bad or half-filled coordinate is still rejected',
          any(r['id'] == 'place-b-c' and 'latitude/longitude' in r['reason'] for r in rep['rejects']))

    p, rep = run(x1, hist)
    check('Sheet without the Place ID columns still loads and says the columns are missing',
          p.returncode == 0 and rep['missingOptionalMovementColumns'] == ['Destination place ID', 'Origin place ID'] and rep['movementsWithArrowEnds'] == 0, rep['missingOptionalMovementColumns'])

    x4 = td / 'four.xlsx'
    wb = openpyxl.load_workbook(x1); ws = wb['Communities']
    ws.cell(1, ws.max_column + 1, 'Concentration'); ws.cell(2, ws.max_column, 'major'); wb.save(x4)
    p, rep = run(x4, hist)
    comm = {c['id']: c for c in json.loads(hist.read_text())['sheet']['communities']}
    check('engine columns load when the Sheet gains them', comm['POP-001'].get('engine') == {'concentration': 'major'}, comm['POP-001'])

    x5 = td / 'five.xlsx'; wb = openpyxl.load_workbook(x1); wb['Places'].delete_cols(8, 1); wb['Places'].delete_cols(5, 1); wb.save(x5)
    p, _ = run(x5, hist)
    check('missing expected column stops the load with a clear message', p.returncode != 0 and 'missing expected columns' in (p.stdout + p.stderr))

print('\n%s' % ('ALL PASS' if not fails else 'FAILED: ' + '; '.join(fails)))
sys.exit(1 if fails else 0)

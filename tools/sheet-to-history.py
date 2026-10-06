#!/usr/bin/env python3
"""Load the Jewish Continuity master dataset (Google Sheet export) into data/history.json.

SOURCE OF TRUTH: the Google Sheet "Jewish_Continuity_Master_Dataset_v0.2.2" (linked in the project
instructions). The build sandbox cannot reach Google, so this script reads an .xlsx export of that
Sheet. Export it fresh each time (File > Download > .xlsx, or ask the dataset chat's Drive connection).
Never commit the export: a committed copy would be an old copy.

USAGE
    pip install openpyxl
    python3 tools/sheet-to-history.py EXPORT.xlsx                 # preview: loads ALL cards, drafts included
    python3 tools/sheet-to-history.py EXPORT.xlsx --dry-run       # report only, writes nothing
    python3 tools/sheet-to-history.py EXPORT.xlsx --public --out data/history.public.json
                                                                  # public build: approved cards only
    options: --history PATH (default data/history.json)  --report PATH (write the full report as JSON)
             --sheet-modified ISO (Sheet's last-modified time, recorded in meta)

WHAT IT DOES TO history.json
    It replaces ONE top-level key, "sheet". Every other key (the engine's places, populations,
    movements, footprints, cameras, eras ...) is left byte for byte as it was.

WHAT GOES IN (public fields only)
    IDs, dates, place IDs, card type, card date, card title, card description, Meanwhile line,
    map symbol, arrow treatment, and the record-to-record links (IDs only).
WHAT NEVER GOES IN
    Columns labelled "never on cards", sources, evidence category, confidence, population figures,
    internal notes, Display Priority notes, the "Map treatment" free text, "Note for Jeffrey".

EVERY RUN REPORTS what was added, changed and removed (by permanent ID), what the validator
rejected (with the reason), and warnings. Nothing is silently fixed. Rejected records are left out
of the output and listed so they can be sent back to the dataset chat.
"""
import argparse, copy, datetime, hashlib, json, re, sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit('openpyxl is required: pip install openpyxl')

SCHEMA = 'sheet-1'
CARD_TYPES = {'Event', 'Community', 'Movement', 'Archaeology'}
CARD_STATUSES = {'approved', 'draft'}
ID_PATTERNS = {
    'EVT': r'EVT-\d{4}', 'POP': r'POP-\d{3}', 'CHG': r'CHG-\d{3}', 'DES': r'DES-\d{3}', 'LOC': r'LOC-\d{3}',
    'MOV': r'MOV-\d{3}', 'CTX': r'CTX-\d{3}', 'VOY': r'VOY-\d{3}', 'EVD': r'EVD-\d{3}',
}
LINK_RE = re.compile(r'\b(?:EVT|POP|CHG|DES|LOC|MOV|CTX|VOY|EVD)-\d+\b')
SYMBOL_KINDS = [  # only the locked/leaning shapes get a kind; everything else stays as the Sheet's own words
    (re.compile(r'^Red octagon', re.I), 'octagon', 'locked'),
    (re.compile(r'^Split diamond', re.I), 'splitDiamond', 'leaning'),
    (re.compile(r'^Purple diamond', re.I), 'purpleDiamond', 'leaning'),
    (re.compile(r'^Black square$', re.I), 'blackSquare', 'locked'),  # approved by Jeffrey Oct 5: community no longer exists
    (re.compile(r'^Candidate:\s*black square', re.I), 'blackSquare', 'candidate'),
]

# Sheet columns, by header text, so a reordered or widened Sheet still loads.
TABS = {
    'Events': dict(id='Event ID', start='Start year', end='End year', dateDisplay='Date (display)',
                   placeId='Place ID', movementIds='Movement record IDs', evidenceIds='Evidence object IDs',
                   cardType='Card type', cardDate='Card date', cardTitle='Card title',
                   cardDescription='Card description', meanwhile='Meanwhile (card line)',
                   meanwhileStatus='Meanwhile status', mapSymbol='Map symbol', cardStatus='Card status'),
    'Communities': dict(id='Record ID', family='Family', dateDisplay='Date (display)', start='Start year',
                        end='End year', placeId='Place ID', mapSymbol='Map symbol', links='Linked records',
                        cardType='Card type', cardDate='Card date', cardTitle='Card title',
                        cardDescription='Card description', cardStatus='Card status'),
    'Movements': dict(id='Movement ID', eventId='Event ID', dateDisplay='Date (display)', start='Start year',
                      end='End year', origin='Origin (historical)', destination='Destination (historical)',
                      arrow='Arrow treatment', links='Linked records'),
    'Context & Voyages': dict(id='Record ID', family='Family', dateDisplay='Date (display)', start='Start year',
                              links='Linked records'),
    'Evidence': dict(id='Evidence ID', start='Creation start yr', end='Creation end yr', links='Linked event IDs'),
    'Places': dict(id='Place ID', historicalName='Historical name', modernName='Modern equivalent', role='Role',
                   lat='Latitude', lon='Longitude', landOfIsrael='Land of Israel?'),
}
# Optional engine-field columns. If the dataset chat adds them to Communities, they load; until then every
# new POP/CHG record is reported as "no engine fields" and nothing is guessed.
OPTIONAL_COMMUNITY_COLS = dict(concentration='Concentration', presenceStatus='Presence status',
                               footprintRadiusKm='Footprint radius (km)')
# Optional arrow-endpoint columns on Movements. The Sheet gives origin and destination as text only; an arrow can be
# drawn only when both ends are Place IDs from the Places tab. Blank = no arrow endpoint (nothing is guessed).
OPTIONAL_MOVEMENT_COLS = dict(originPlaceId='Origin place ID', destinationPlaceId='Destination place ID')
SECTION_OF_TAB = {'Events': 'events', 'Communities': 'communities', 'Movements': 'movements',
                  'Context & Voyages': 'context', 'Evidence': 'evidence', 'Places': 'places'}

# Era and camera tabs (added with master v0.2.2). They are optional as a whole (an older export has none of them), but a tab
# that is present must carry every expected column: a damaged tab stops the load, it is never loosened to fit.
ERA_TABS = {
    'Eras': dict(era='Era', name='Era name (proposed)', start='Start year', startDisplay='Start (display)',
                 lastYear='Last year inside era', summary='Summary line (one sentence, proposed)',
                 centerPlaceId='Opening camera centre (Place ID)', frame='Frame must keep on screen (Place IDs)',
                 status='Status'),
    'Camera Stops': dict(stop='Stop', era='Era', start='Start year', label='Stop label (proposed)',
                         centerPlaceId='Camera centre (Place ID)', role='Centre scale (Places Role)',
                         frame='Frame must keep on screen (Place IDs)', records='Justifying Sheet records (IDs)',
                         status='Status'),
    'Era Assignment': dict(id='Record ID', era='Era (recommended boundaries)', runsPast='Runs past era end?'),
}
# Camera zoom rule (approved by Jeffrey Oct 6, 2026). The Sheet has no zoom numbers, only a centre Place and the Places that
# must stay on screen. half-width (degrees of arc, the engine's own unit) = distance from the centre to the farthest frame
# place x 1.15, never below a floor set by the centre's Role (settlement 12, region 24), clamped to the engine's 10..45.
VIEW_MARGIN, VIEW_FLOOR, VIEW_MIN, VIEW_MAX = 1.15, {'settlement': 12.0, 'region': 24.0}, 10.0, 45.0


def clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def read_tab(wb, tab, cols, optional=None):
    ws = wb[tab]
    header = [c.value for c in ws[1]]
    pos = {h: i for i, h in enumerate(header) if h}
    missing = [h for h in cols.values() if h not in pos]
    if missing:
        raise SystemExit(f'Tab "{tab}" is missing expected columns: {missing}')
    use = dict(cols)
    for k, h in (optional or {}).items():
        if h in pos:
            use[k] = h
    rows = []
    for rnum, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not r or r[0] in (None, ''):
            continue
        rec = {k: clean(r[pos[h]]) if pos[h] < len(r) else None for k, h in use.items()}
        rec['_row'] = rnum
        rows.append(rec)
    return rows, sorted(set((optional or {}).values()) - set(pos))


def links(s):
    return sorted(set(LINK_RE.findall(str(s or ''))))


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def symbol(raw):
    if not raw:
        return None
    for rx, kind, status in SYMBOL_KINDS:
        if rx.match(raw):
            return {'text': raw, 'kind': kind, 'status': status}
    return {'text': raw, 'kind': None, 'status': None}


class Report:
    def __init__(self):
        self.rejects, self.warnings, self.withheld = [], [], []

    def reject(self, tab, rid, row, why):
        self.rejects.append({'tab': tab, 'id': rid, 'row': row, 'reason': why})

    def warn(self, tab, rid, row, why):
        self.warnings.append({'tab': tab, 'id': rid, 'row': row, 'note': why})


def arc_deg(a, b):
    """Great-circle distance in degrees between two (lat, lon) points."""
    import math
    p1, l1, p2, l2 = map(math.radians, (a[0], a[1], b[0], b[1]))
    c = math.sin(p1) * math.sin(p2) + math.cos(p1) * math.cos(p2) * math.cos(l1 - l2)
    return math.degrees(math.acos(max(-1.0, min(1.0, c))))


def read_optional_tab(wb, tab, cols):
    """Rows of an era/camera tab, or None when the workbook has no such tab. Present but damaged = stop."""
    if tab not in wb.sheetnames:
        return None
    rows, _ = read_tab(wb, tab, cols)
    return rows


def build_eras(wb, places, all_ids, rep):
    """Eras, Camera Stops and Era Assignment -> three sections. Nothing is guessed: a bad row is rejected with its reason."""
    by_id = {p['id']: p for p in places if p['lat'] is not None}
    raw = {t: read_optional_tab(wb, t, c) for t, c in ERA_TABS.items()}
    for t, v in raw.items():
        if v is None:
            rep.warn(t, '-', 0, f'tab "{t}" is not in this export; no {t.lower()} loaded')
    status_of = lambda v: (str(v).strip().lower() if v else '')

    def frame_ids(tab, r, label):
        txt = str(r.get('frame') or '').strip()
        if not txt or txt.startswith('('):  # "(this place only)" = no extra frame places
            return []
        ids = [x.strip() for x in txt.split(';') if x.strip()]
        bad = [i for i in ids if i not in by_id]
        if bad:
            rep.reject(tab, label, r['_row'], f'frame place IDs not in Places (or without coordinates): {bad}'); return None
        return ids

    eras = []
    for r in raw['Eras'] or []:
        n, label = r.get('era'), f"ERA-{r.get('era')}"
        if not isinstance(n, int) or not 1 <= n <= 7:
            rep.reject('Eras', label, r['_row'], f'era number is not 1 to 7: {n!r}'); continue
        if num(r.get('start')) is None or not r.get('name'):
            rep.reject('Eras', label, r['_row'], 'era needs a name and a numeric start year'); continue
        st = status_of(r.get('status'))
        reference = st.startswith('reference only')
        if not reference and st not in ('approved', 'draft'):
            rep.reject('Eras', label, r['_row'], f'status {r.get("status")!r} is not approved, draft or reference only'); continue
        e = {'id': label, 'era': n, 'name': r['name'], 'startYear': r['start'], 'startDisplay': r.get('startDisplay'),
             'lastYear': r.get('lastYear'), 'summary': r.get('summary'), 'status': 'reference' if reference else st}
        if not reference:
            c = r.get('centerPlaceId')
            if c not in by_id:
                rep.reject('Eras', label, r['_row'], f'opening camera centre {c!r} is not a Place with coordinates'); continue
            fr = frame_ids('Eras', r, label)
            if fr is None: continue
            e.update(centerPlaceId=c, framePlaceIds=fr)
        eras.append(e)
    eras.sort(key=lambda e: e['era'])
    ok_eras = {e['era'] for e in eras}

    stops = []
    for r in raw['Camera Stops'] or []:
        sid = str(r.get('stop')) if r.get('stop') is not None else None
        label = f'STOP-{sid}'
        if not sid or not re.fullmatch(r'[4-7]\.\d{1,2}', sid):
            rep.reject('Camera Stops', label, r['_row'], 'stop number is not like 4.1 (era 4 to 7)'); continue
        if r.get('era') != int(sid.split('.')[0]) or r.get('era') not in ok_eras:
            rep.reject('Camera Stops', label, r['_row'], f'era {r.get("era")!r} does not match the stop number or is not a loaded era'); continue
        if num(r.get('start')) is None:
            rep.reject('Camera Stops', label, r['_row'], f'start year is not a number: {r.get("start")!r}'); continue
        st = status_of(r.get('status'))
        if st not in ('approved', 'draft'):
            rep.reject('Camera Stops', label, r['_row'], f'status {r.get("status")!r} is not approved or draft'); continue
        c, role = r.get('centerPlaceId'), status_of(r.get('role'))
        if c not in by_id:
            rep.reject('Camera Stops', label, r['_row'], f'camera centre {c!r} is not a Place with coordinates'); continue
        if role not in VIEW_FLOOR:
            rep.reject('Camera Stops', label, r['_row'], f'centre scale {r.get("role")!r} is not settlement or region'); continue
        fr = frame_ids('Camera Stops', r, label)
        if fr is None: continue
        reach = max([arc_deg((by_id[c]['lat'], by_id[c]['lon']), (by_id[i]['lat'], by_id[i]['lon'])) for i in fr] + [0.0])
        want = max(reach * VIEW_MARGIN, VIEW_FLOOR[role])
        fits = want <= VIEW_MAX
        if not fits:
            rep.warn('Camera Stops', label, r['_row'],
                     f'the frame reaches {reach:.0f} degrees from the centre, so the view would need {want:.0f} degrees of half-width; the engine shows at most {VIEW_MAX:.0f}, so the view is clamped and part of the frame is off screen')
        stops.append({'id': label, 'stop': sid, 'era': r['era'], 'startYear': r['start'], 'label': r.get('label'),
                      'centerPlaceId': c, 'centerLat': by_id[c]['lat'], 'centerLon': by_id[c]['lon'], 'role': role,
                      'framePlaceIds': fr, 'reachDeg': round(reach, 1), 'viewDeg': round(min(VIEW_MAX, max(VIEW_MIN, want)), 1),
                      'fits': fits, 'recordIds': links(r.get('records')), 'status': st})
    stops.sort(key=lambda x: (x['startYear'], [int(p) for p in x['stop'].split('.')]))
    for a, b in zip(stops, stops[1:]):
        if a['startYear'] == b['startYear']:
            rep.warn('Camera Stops', b['id'], 0, f'starts in the same year as {a["id"]} ({a["startYear"]}); by year the camera shows {b["id"]}, and cards in {a["id"]} turn the globe on their own')

    assign = []
    for r in raw['Era Assignment'] or []:
        rid, e = r.get('id'), str(r.get('era') or '').strip()
        m = re.fullmatch(r'Era ([1-7])', e)
        if not m and not e.startswith('After 1897'):
            rep.reject('Era Assignment', rid, r['_row'], f'era {e!r} is not Era 1 to 7 or After 1897'); continue
        if rid not in all_ids:
            rep.warn('Era Assignment', rid, r['_row'], 'record ID is not in the Sheet tabs this loader reads')
        assign.append({'id': rid, 'era': int(m.group(1)) if m else None, 'runsPast': r.get('runsPast')})
    return eras, stops, assign


def build(wb, rep):
    raw, missing_optional = {}, {'Communities': [], 'Movements': []}
    for tab, cols in TABS.items():
        opt = {'Communities': OPTIONAL_COMMUNITY_COLS, 'Movements': OPTIONAL_MOVEMENT_COLS}.get(tab)
        raw[tab], miss = read_tab(wb, tab, cols, opt)
        if tab in missing_optional:
            missing_optional[tab] = miss

    seen = {}
    all_ids = set()
    for tab, rows in raw.items():  # pass 1: ID format + uniqueness
        keep = []
        for r in rows:
            rid = r['id']
            m = re.match(r'[A-Za-z]+', str(rid)) if rid else None
            fam = m.group(0) if m else None
            if rid is not None and len(str(rid)) > 40:
                rid = str(rid)[:37] + '...'  # stray pasted text in the ID column: shorten for the report
            if tab == 'Places':
                ok = bool(re.fullmatch(r'place-[A-Za-z0-9]+(?:-[A-Za-z0-9]+)*', str(rid)))
            else:
                ok = fam in ID_PATTERNS and re.fullmatch(ID_PATTERNS[fam], str(rid))
            if not ok:
                rep.reject(tab, rid, r['_row'], 'ID is not in a permanent-ID format'); continue
            if (tab, rid) in seen:
                rep.reject(tab, rid, r['_row'], f'duplicate ID (first on row {seen[(tab, rid)]})'); continue
            seen[(tab, rid)] = r['_row']
            keep.append(r)
            if tab != 'Places':
                all_ids.add(rid)
        raw[tab] = keep

    # places
    places, place_ids, placed = [], set(), set()
    for r in raw['Places']:
        no_coords = r['lat'] in (None, '') and r['lon'] in (None, '')
        if no_coords:
            r['lat'] = r['lon'] = None
            rep.warn('Places', r['id'], r['_row'], 'no latitude/longitude yet: the place loads, but nothing can be drawn there until coordinates are supplied')
        elif num(r['lat']) is None or num(r['lon']) is None or not (-90 <= r['lat'] <= 90 and -180 <= r['lon'] <= 180):
            rep.reject('Places', r['id'], r['_row'], 'latitude/longitude missing or out of range'); continue
        place_ids.add(r['id'])
        if r['lat'] is not None: placed.add(r['id'])
        places.append({'id': r['id'], 'historicalName': r['historicalName'], 'modernName': r['modernName'],
                       'role': r['role'], 'lat': r['lat'], 'lon': r['lon'],
                       'landOfIsrael': str(r['landOfIsrael'] or '').lower() == 'yes'})

    def years(tab, r, need_start=True):
        s, e = r.get('start'), r.get('end')
        if r.get('start') is not None and num(s) is None:
            rep.reject(tab, r['id'], r['_row'], f'start year is not a number: {s!r}'); return None
        if need_start and s is None:
            rep.reject(tab, r['id'], r['_row'], 'start year is empty'); return None
        if e is not None and num(e) is None:
            rep.reject(tab, r['id'], r['_row'], f'end year is not a number: {e!r}'); return None
        if e is not None and s is not None and e < s:
            rep.reject(tab, r['id'], r['_row'], f'end year {e} is before start year {s}'); return None
        return s, e

    def check_links(tab, r, ids):
        bad = [i for i in ids if i not in all_ids]
        if bad:
            rep.warn(tab, r['id'], r['_row'], f'links to IDs that are not in the Sheet: {bad}')
        return [i for i in ids if i in all_ids]

    cards, card_text_seen = [], {}

    def make_card(tab, r, rid_extra=None):
        title = r.get('cardTitle')
        if not any(r.get(k) for k in ('cardType', 'cardDate', 'cardTitle', 'cardDescription', 'cardStatus')):
            return None
        missing = [n for n, k in (('card type', 'cardType'), ('card date', 'cardDate'), ('card title', 'cardTitle'),
                                  ('card description', 'cardDescription'), ('card status', 'cardStatus')) if not r.get(k)]
        if missing:
            rep.reject(tab + ' (card only)', r['id'], r['_row'],
                       f'card is incomplete, so no card is made (record itself is kept): missing {", ".join(missing)}')
            return None
        if r['cardType'] not in CARD_TYPES:
            rep.reject(tab + ' (card only)', r['id'], r['_row'], f'card type {r["cardType"]!r} is not one of {sorted(CARD_TYPES)}'); return None
        st = str(r['cardStatus']).lower()
        if st not in CARD_STATUSES:
            rep.reject(tab + ' (card only)', r['id'], r['_row'], f'card status {r["cardStatus"]!r} is not approved/draft'); return None
        key = (r['cardDate'], r['cardTitle'], r['cardDescription'])
        if key in card_text_seen:
            rep.warn(tab, r['id'], r['_row'], f'card text is identical to {card_text_seen[key]}; no second card made (the {card_text_seen[key]} card is used)')
            return None
        card_text_seen[key] = r['id']
        card = {'id': r['id'], 'recordId': r['id'], 'type': r['cardType'], 'date': r['cardDate'], 'title': title,
                'description': r['cardDescription'], 'status': st, 'startYear': r.get('start'), 'endYear': r.get('end'),
                'placeId': r.get('placeId')}
        if r.get('meanwhile'):
            if r['cardType'] == 'Event':
                card['meanwhile'] = r['meanwhile']
                card['meanwhileStatus'] = str(r.get('meanwhileStatus') or 'unstated').lower()
            else:  # Jeffrey, Oct 5: Meanwhile shows on Event cards only. The line stays in the Sheet; it is not exported.
                rep.withheld.append({'id': r['id'], 'row': r['_row'], 'cardType': r['cardType']})
        if r.get('movementIds'):
            card['movementIds'] = r['movementIds']
        if r.get('evidenceIds'):
            card['evidenceIds'] = r['evidenceIds']
        sym = symbol(r.get('mapSymbol'))
        if sym:
            card['mapSymbol'] = sym
        return card

    # events (first, so an Events card wins over an identical Communities card)
    events = []
    for r in raw['Events']:
        y = years('Events', r)
        if y is None: continue
        if r['placeId'] and r['placeId'] not in place_ids:
            rep.reject('Events', r['id'], r['_row'], f'Place ID {r["placeId"]!r} is not in the Places tab'); continue
        r['movementIds'] = check_links('Events', r, links(r.get('movementIds')))
        r['evidenceIds'] = check_links('Events', r, links(r.get('evidenceIds')))
        card = make_card('Events', r)
        if card: cards.append(card)
        rec = {'id': r['id'], 'startYear': y[0], 'endYear': y[1], 'dateDisplay': r['dateDisplay'], 'placeId': r['placeId'],
               'movementIds': r['movementIds'], 'evidenceIds': r['evidenceIds'], 'hasCard': bool(card)}
        sym = symbol(r.get('mapSymbol'))
        if sym: rec['mapSymbol'] = sym
        events.append(rec)

    communities = []
    for r in raw['Communities']:
        y = years('Communities', r)
        if y is None: continue
        if not r['placeId']:
            rep.reject('Communities', r['id'], r['_row'], 'no Place ID, so the record cannot be placed on the map'); continue
        if r['placeId'] not in place_ids:
            rep.reject('Communities', r['id'], r['_row'], f'Place ID {r["placeId"]!r} is not in the Places tab'); continue
        lk = check_links('Communities', r, links(r.get('links')))
        card = make_card('Communities', r)
        if card: cards.append(card)
        rec = {'id': r['id'], 'family': r['family'], 'startYear': y[0], 'endYear': y[1], 'dateDisplay': r['dateDisplay'],
               'placeId': r['placeId'], 'links': lk, 'hasCard': bool(card)}
        sym = symbol(r.get('mapSymbol'))
        if sym:
            rec['mapSymbol'] = sym
            if sym['status'] == 'candidate':
                rep.warn('Communities', r['id'], r['_row'], 'symbol is a candidate whose data rule is still OPEN; not to be drawn yet')
        for k in OPTIONAL_COMMUNITY_COLS:
            if r.get(k) is not None: rec.setdefault('engine', {})[k] = r[k]
        if r['family'] in ('POP', 'CHG') and 'engine' not in rec:
            rec['noEngineFields'] = True
        communities.append(rec)

    movements = []
    for r in raw['Movements']:
        y = years('Movements', r)
        if y is None: continue
        if r['eventId'] and r['eventId'] not in all_ids:
            rep.warn('Movements', r['id'], r['_row'], f'Event ID {r["eventId"]!r} is not in the Events tab')
        lk = check_links('Movements', r, links(r.get('links')))
        arrow = r['arrow']
        if arrow and 'not stated' in arrow.lower():
            rep.warn('Movements', r['id'], r['_row'], 'arrow treatment is "not stated"; no colour can be chosen yet')
        ends, bad_end = {}, False
        for key, label in (('originPlaceId', 'origin'), ('destinationPlaceId', 'destination')):
            v = r.get(key)
            if v is not None and v not in place_ids:
                rep.reject('Movements', r['id'], r['_row'], f'{label} place ID {v!r} is not in the Places tab'); bad_end = True
            ends[key] = v
        if bad_end: continue
        if (ends['originPlaceId'] is None) != (ends['destinationPlaceId'] is None) and (
                'originPlaceId' in r and 'destinationPlaceId' in r):
            rep.warn('Movements', r['id'], r['_row'], 'only one end has a Place ID (general location only), so no arrow can be drawn')
        movements.append({'id': r['id'], 'eventId': r['eventId'], 'startYear': y[0], 'endYear': y[1],
                          'dateDisplay': r['dateDisplay'], 'origin': r['origin'], 'destination': r['destination'],
                          'originPlaceId': ends['originPlaceId'], 'destinationPlaceId': ends['destinationPlaceId'],
                          'arrowEndsReady': ends['originPlaceId'] is not None and ends['destinationPlaceId'] is not None
                                            and ends['originPlaceId'] in placed and ends['destinationPlaceId'] in placed,
                          'arrowTreatment': arrow, 'links': lk})

    context = []
    for r in raw['Context & Voyages']:
        y = years('Context & Voyages', r)
        if y is None: continue
        context.append({'id': r['id'], 'family': r['family'], 'startYear': y[0], 'dateDisplay': r['dateDisplay'],
                        'links': check_links('Context & Voyages', r, links(r.get('links')))})

    evidence = []
    for r in raw['Evidence']:
        y = years('Evidence', r)
        if y is None: continue
        evidence.append({'id': r['id'], 'startYear': y[0], 'endYear': y[1],
                         'links': check_links('Evidence', r, links(r.get('links')))})

    cards.sort(key=lambda c: (c['startYear'], c['id']))
    sections = {'places': places, 'events': events, 'communities': communities, 'movements': movements,
                'context': context, 'evidence': evidence, 'cards': cards}
    eras, stops, assign = build_eras(wb, places, all_ids, rep)
    if eras or stops or assign:
        sections.update({'eras': eras, 'cameraStops': stops, 'eraAssignment': assign})

    return sections, missing_optional


def public_view(sections):
    """Public build: approved cards only; Meanwhile only if its own status is approved."""
    out = copy.deepcopy(sections)
    out['cards'] = [c for c in out['cards'] if c['status'] == 'approved']
    keep = {c['id'] for c in out['cards']}
    for c in out['cards']:
        if c.get('meanwhile') and c.get('meanwhileStatus') != 'approved':
            c.pop('meanwhile'); c.pop('meanwhileStatus')
    for sec in ('events', 'communities'):
        for r in out[sec]:
            r['hasCard'] = r['id'] in keep
    return out


def diff(old, new):
    res = {}
    for sec in sorted(set(old) | set(new)):
        o = {r['id']: r for r in old.get(sec, [])}
        n = {r['id']: r for r in new.get(sec, [])}
        changed = sorted(i for i in set(o) & set(n) if o[i] != n[i])
        res[sec] = {'added': sorted(set(n) - set(o)), 'removed': sorted(set(o) - set(n)), 'changed': changed,
                    'changedFields': {i: sorted(k for k in set(o[i]) | set(n[i]) if o[i].get(k) != n[i].get(k)) for i in changed}}
    return res


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx')
    ap.add_argument('--history', default='data/history.json')
    ap.add_argument('--out')
    ap.add_argument('--public', action='store_true')
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--report')
    ap.add_argument('--sheet-modified')
    a = ap.parse_args(argv)

    xlsx = Path(a.xlsx)
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    rep = Report()
    sections, missing_optional = build(wb, rep)
    if a.public:
        sections = public_view(sections)

    hist_path = Path(a.history)
    text = hist_path.read_text(encoding='utf-8')
    hist = json.loads(text)
    old = {k: v for k, v in (hist.get('sheet') or {}).items() if isinstance(v, list)}
    new = {k: v for k, v in sections.items()}
    d = diff(old, new)

    meta = {'schema': SCHEMA, 'source': 'Google Sheet Jewish_Continuity_Master_Dataset_v0.2.2',
            'sheetFileId': '1XbmQI73C3Ik-vePL-WO5_lgHWMFWTYST1hCZLCyGbV8',
            'sheetModified': a.sheet_modified or (hist.get('sheet') or {}).get('meta', {}).get('sheetModified'),
            'exportSha256': hashlib.sha256(xlsx.read_bytes()).hexdigest()[:16],
            'mode': 'public' if a.public else 'preview',
            'counts': {k: len(v) for k, v in sections.items()}}
    changed_any = any(v['added'] or v['removed'] or v['changed'] for v in d.values())
    prev_mode = (hist.get('sheet') or {}).get('meta', {}).get('mode')
    # Only touch the file when records changed (or the mode changed); a re-run of the same export is a no-op.
    out_hist = dict(hist)
    if changed_any or prev_mode != meta['mode'] or 'sheet' not in hist:
        out_hist['sheet'] = {'meta': meta, **sections}
    else:
        out_hist['sheet'] = hist['sheet']

    report = {'mode': meta['mode'], 'counts': meta['counts'], 'diff': d, 'rejects': rep.rejects,
              'warnings': rep.warnings, 'meanwhileWithheldNonEvent': rep.withheld,
              'missingOptionalEngineColumns': missing_optional['Communities'],
              'missingOptionalMovementColumns': missing_optional['Movements'],
              'movementsWithArrowEnds': sum(1 for m in sections['movements'] if m['arrowEndsReady']),
              'recordsWithNoEngineFields': sum(1 for c in sections['communities'] if c.get('noEngineFields'))}

    # ---- console report
    print(f'Loaded {xlsx.name}  mode={meta["mode"]}')
    print('Counts: ' + ', '.join(f'{k} {v}' for k, v in meta['counts'].items()))
    print('\nChanges since the last load (by permanent ID):')
    for sec, v in d.items():
        print(f'  {sec:12} added {len(v["added"]):3}  changed {len(v["changed"]):3}  removed {len(v["removed"]):3}')
        for k in ('added', 'changed', 'removed'):
            if 0 < len(v[k]) <= 8: print(f'      {k}: {", ".join(v[k])}')
        if 8 < len(v['changed']) <= 60 or (v['changed'] and len(v['changed']) <= 8):
            for i in v['changed']: print(f'        {i}: {", ".join(v["changedFields"][i])}')
    print(f'\nValidator rejects: {len(rep.rejects)}')
    for x in rep.rejects:
        print(f'  REJECT {x["tab"]} row {x["row"]} {x["id"]}: {x["reason"]}')
    print(f'Warnings: {len(rep.warnings)}')
    for x in rep.warnings:
        print(f'  warn   {x["tab"]} row {x["row"]} {x["id"]}: {x["note"]}')
    print(f'Meanwhile lines withheld because the card is not an Event card: {len(rep.withheld)}')
    for x in rep.withheld:
        print(f'  held   Events row {x["row"]} {x["id"]} ({x["cardType"]} card)')
    if missing_optional['Communities']:
        print(f'\nCommunities tab has no column for: {missing_optional["Communities"]}')
    if missing_optional['Movements']:
        print(f'Movements tab has no column for: {missing_optional["Movements"]}')
    print(f'Movements with both arrow ends as Place IDs (can be drawn): {report["movementsWithArrowEnds"]} of {len(sections["movements"])}')
    print(f'POP/CHG records with no engine fields (no shading until supplied): {report["recordsWithNoEngineFields"]}')

    if a.report:
        Path(a.report).write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding='utf-8')
    if a.dry_run:
        print('\nDry run: nothing written.')
        return 0
    out_path = Path(a.out) if a.out else hist_path
    new_text = json.dumps(out_hist, indent=2, ensure_ascii=False)
    if out_path == hist_path and new_text == text:
        print('\nNo changes: data/history.json already matches this export.')
    else:
        out_path.write_text(new_text, encoding='utf-8')
        print(f'\nWrote {out_path}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
